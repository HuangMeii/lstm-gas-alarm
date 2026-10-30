from pathlib import Path
import re

import pandas as pd
from openpyxl import load_workbook

from .config import NIST_RAW_DIR, PROCESSED_DIR, SUMMARY_WORKBOOK


def _clean_columns(columns: list[str]) -> list[str]:
    return [re.sub(r"\s+", " ", str(column).strip()) for column in columns]


def read_smokebox_file(path: Path) -> pd.DataFrame:
    frame = pd.read_csv(path, sep="\t", skiprows=1)
    frame.columns = _clean_columns(frame.columns.tolist())
    frame["experiment_id"] = path.stem
    frame["experiment_type"] = "smokebox"
    frame["alarm_time_sec"] = pd.to_numeric(
        frame.get("Alarm?"), errors="coerce"
    ).where(lambda values: values == 1).groupby(frame.index).transform("max")
    alarm_indices = frame.index[frame["Alarm?"].eq(1)]
    if len(alarm_indices):
        frame["alarm_time_sec"] = frame.loc[alarm_indices[0], "Counter"]
    else:
        frame["alarm_time_sec"] = pd.NA
    return frame


def room_alarm_manifest() -> pd.DataFrame:
    workbook = load_workbook(SUMMARY_WORKBOOK, data_only=True, read_only=True)
    records: list[dict[str, object]] = []
    for sheet in workbook.sheetnames:
        if not sheet.startswith("Room "):
            continue
        rows = list(workbook[sheet].iter_rows(values_only=True))
        for column in range(len(rows[3])):
            prefix = rows[3][column]
            if not isinstance(prefix, str) or "#" not in prefix:
                continue
            for row in range(4, len(rows), 4):
                model_unit = rows[row][column - 2] if column >= 2 else None
                if not isinstance(model_unit, str):
                    continue
                for trial in range(3):
                    alarm = rows[row + 1][column + trial]
                    trial_code = rows[row][column + trial]
                    if isinstance(alarm, (int, float)) and isinstance(trial_code, str):
                        group, repetition = trial_code.split(",", maxsplit=1)
                        filename = prefix.replace("#", group, 1).replace(
                            "#", repetition, 1
                        )
                        records.append(
                            {
                                "experiment_id": filename,
                                "alarm_time_sec": float(alarm),
                                "experiment_type": "room",
                                "alarm_model_unit": model_unit,
                            }
                        )
    return pd.DataFrame(records).drop_duplicates("experiment_id")


def read_room_file(path: Path, manifest: pd.DataFrame) -> pd.DataFrame:
    frame = pd.read_csv(path, sep="\t")
    frame.columns = _clean_columns(frame.columns.tolist())
    frame["experiment_id"] = path.stem
    frame["experiment_type"] = "room"
    match = manifest[manifest["experiment_id"].eq(path.stem)]
    frame["alarm_time_sec"] = (
        float(match.iloc[0]["alarm_time_sec"]) if not match.empty else pd.NA
    )
    return frame


def build_manifest() -> pd.DataFrame:
    manifest = room_alarm_manifest()
    for path in (NIST_RAW_DIR / "smokebox").rglob("*.tsv"):
        frame = read_smokebox_file(path)
        alarm = frame["alarm_time_sec"].dropna()
        manifest = pd.concat(
            [
                manifest,
                pd.DataFrame(
                    [
                        {
                            "experiment_id": path.stem,
                            "alarm_time_sec": float(alarm.iloc[0])
                            if not alarm.empty
                            else None,
                            "experiment_type": "smokebox",
                        }
                    ]
                ),
            ],
            ignore_index=True,
        )
    return manifest.sort_values(["experiment_id", "alarm_model_unit"])


def prepare_nist() -> pd.DataFrame:
    PROCESSED_DIR.mkdir(parents=True, exist_ok=True)
    manifest = build_manifest()
    manifest.to_csv(PROCESSED_DIR / "experiment_manifest.csv", index=False)
    return manifest
